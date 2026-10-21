"""Generate the .pptx fixture used by the integration tests.

Run from the project root:
    python3 test/fixtures/make_fixture.py
"""

import os

from pptx import Presentation
from pptx.util import Inches


def main():
    prs = Presentation()

    # Slide 1: title + body placeholder.
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Quarterly Report"
    slide.placeholders[1].text = "Revenue grew 12%\nCosts flat"

    # Slide 2: table, grouped shape, notes.
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    rows, cols = 2, 2
    table = slide.shapes.add_table(
        rows, cols, Inches(1), Inches(1), Inches(4), Inches(1)
    ).table
    table.cell(0, 0).text = "Region"
    table.cell(0, 1).text = "Sales"
    table.cell(1, 0).text = "EMEA"
    table.cell(1, 1).text = "42"

    group = slide.shapes.add_group_shape()
    group.shapes.add_textbox(Inches(1), Inches(3), Inches(3), Inches(1)).text_frame.text = (
        "Grouped caption"
    )

    slide.shapes.add_textbox(Inches(1), Inches(4), Inches(3), Inches(1)).text_frame.text = (
        "Appendix notes"
    )
    slide.notes_slide.notes_text_frame.text = "Speaker note for slide two"

    output = os.path.join(os.path.dirname(__file__), "sample.pptx")
    prs.save(output)
    print(f"wrote {output}")


if __name__ == "__main__":
    main()
